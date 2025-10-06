import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
import openpyxl
import os

excel_file = "expenses.xlsx"

# --- Excel setup ---
if not os.path.exists(excel_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expenses"
    ws.append(["ID", "Date", "Category", "Amount", "Description"])
    wb.save(excel_file)

def get_next_id(sheet):
    if sheet.max_row == 1:
        return 1
    ids = [sheet.cell(row=i, column=1).value for i in range(2, sheet.max_row + 1)]
    return max(ids) + 1 if ids else 1

def load_data():
    for i in tree.get_children():
        tree.delete(i)
    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", "end", values=row)
    wb.close()

def clear_entries():
    entry_date.delete(0, tk.END)
    entry_category.set("")
    entry_amount.delete(0, tk.END)
    entry_desc.delete(0, tk.END)

def add_expense():
    date = entry_date.get()
    category = entry_category.get()
    amount = entry_amount.get()
    desc = entry_desc.get()

    if not date or not category or not amount:
        messagebox.showwarning("Input Error", "Date, Category and Amount are required")
        return

    try:
        amount = float(amount)
    except:
        messagebox.showwarning("Input Error", "Amount must be a number")
        return

    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active
    new_id = get_next_id(ws)
    ws.append([new_id, date, category, amount, desc])
    wb.save(excel_file)
    wb.close()

    clear_entries()
    load_data()

def delete_expense():
    selected = tree.selection()
    if not selected:
        return
    item = tree.item(selected[0])
    expense_id = item['values'][0]

    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active

    for i in range(2, ws.max_row + 1):
        if ws.cell(row=i, column=1).value == expense_id:
            ws.delete_rows(i)
            break
    wb.save(excel_file)
    wb.close()

    load_data()

def edit_expense():
    select
